"""แปลง Markdown (เช่น docs/CODE_GUIDE.md) เป็น .docx แบบง่าย

ใช้งาน:  uv run --no-project --with python-docx python docs/report/md_to_docx.py docs/CODE_GUIDE.md
รองรับ: หัวข้อ #–####, รูป ![](path), ย่อหน้า, รายการ -, 1., ตาราง |, code block ```, blockquote >, **ตัวหนา**, `code`, [ลิงก์](url)
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT, MONO = "Leelawadee UI", "Consolas"
GREEN, LEAF, MUTED = RGBColor(0x1F, 0x5E, 0x3B), RGBColor(0x3E, 0x8E, 0x5A), RGBColor(0x66, 0x75, 0x6B)
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def font(run, size=None, bold=None, color=None, mono=False):
    name = MONO if mono else FONT
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT if a == "w:cs" else name)
    if size:
        run.font.size = Pt(size)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(int(size * 2)))
        rpr.append(szcs)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_inline(p, text, size=10, color=None, bold=False):
    text = text.replace("**`", "`").replace("`**", "`")  # ตัวหนา + โค้ดซ้อนกัน → แสดงเป็นโค้ด
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            font(p.add_run(part[2:-2]), size, True, color)
        elif part.startswith("`"):
            font(p.add_run(part[1:-1]), size - 0.5, bold, RGBColor(0x9A, 0x3B, 0x26), mono=True)
        elif part.startswith("["):
            label, target = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            m = re.search(r"#L(\d+)", target)
            font(p.add_run(label), size, True if bold else None, LEAF)
            if m and not label.strip().isdigit() and "บรรทัด" not in label and not re.match(r"^\d+[–-]\d+$", label):
                font(p.add_run(f" (L{m.group(1)})"), size - 1.5, None, MUTED)
        else:
            font(p.add_run(part.replace("\\|", "|")), size, bold or None, color)


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._element.get_or_add_tcPr().append(shd)


def add_table(doc, rows):
    cells = [[c.strip() for c in re.split(r"(?<!\\)\|", r.strip().strip("|"))] for r in rows]
    header, body = cells[0], [r for r in cells[2:]]
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        add_inline(c.paragraphs[0], h, 9, GREEN, bold=True)
        shade(c, "DCEBDF")
    for r, row in enumerate(body, 1):
        for i in range(len(header)):
            c = t.rows[r].cells[i]
            c.text = ""
            add_inline(c.paragraphs[0], row[i] if i < len(row) else "", 9)
    if len(header) == 2:  # คอลัมน์แรกแคบ (เลขบรรทัด / ชื่อไฟล์)
        for row in t.rows:
            row.cells[0].width, row.cells[1].width = Cm(3.2), Cm(12.7)
    doc.add_paragraph()


def convert(src: Path, out: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    lines = src.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.4)
            for j, b in enumerate(block):
                font(p.add_run(b), 8.5, color=RGBColor(0x33, 0x33, 0x33), mono=True)
                if j < len(block) - 1:
                    p.add_run().add_break()
        elif m := re.match(r"^!\[[^\]]*\]\(([^)]+)\)", line):
            doc.add_picture(str(src.parent / m.group(1)), width=Cm(16.6))
            doc.paragraphs[-1].alignment = 1  # กึ่งกลาง
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        elif m := re.match(r"^(#{1,4}) (.*)", line):
            level = len(m.group(1))
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14 if level <= 2 else 8)
            p.paragraph_format.keep_with_next = True
            add_inline(p, m.group(2), {1: 20, 2: 16, 3: 13, 4: 11.5}[level], GREEN if level <= 2 else LEAF, bold=True)
        elif line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            add_inline(p, line.lstrip("> ").strip(), 9.5, MUTED)
        elif re.match(r"^\s*- ", line):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, re.sub(r"^\s*- ", "", line), 10)
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", line), 10)
        elif line.strip() == "---":
            pass
        elif line.strip():
            add_inline(doc.add_paragraph(), line, 10)
        i += 1
    doc.save(out)
    print("saved", out)


if __name__ == "__main__":
    src = Path(sys.argv[1])
    convert(src, src.with_suffix(".docx"))
